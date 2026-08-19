"""Extensión del panel: padrón de convocados + importación Excel/CSV.

Importa ``panel.py`` para conservar la mesa de ayuda existente y agrega rutas
administrativas para cargar el padrón, consultar avance y exportarlo.
"""
import csv
import io
from typing import Any

import psycopg
from fastapi import Depends, File, HTTPException, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from openpyxl import load_workbook

import panel as base

app = base.app

# Añadir navegación sin reescribir el panel de tickets.
base.PAGE = base.PAGE.replace(
    "</style>",
    """
  nav { margin-left:auto; display:flex; gap:14px; flex-wrap:wrap; }
  nav a { color:#fff; font-weight:600; font-size:14px; }
  .ok { color:#15803d; font-weight:700; }
  .warn { color:#b45309; font-weight:700; }
  .danger { color:#b91c1c; font-weight:700; }
  .actions { display:flex; gap:10px; flex-wrap:wrap; margin:0 0 16px; }
  .actions a { background:#0b5cab; color:#fff; padding:9px 13px; border-radius:8px; font-weight:600; }
  .progress { height:10px; background:#e5e7eb; border-radius:999px; overflow:hidden; min-width:110px; }
  .progress span { display:block; height:100%; background:#16a34a; }
</style>""",
).replace(
    "<header><h1>🆘 Mesa de ayuda — Representantes PAN Querétaro</h1></header>",
    """<header><h1>🆘 Acción Electoral — Control y soporte</h1>
<nav><a href=\"/\">Tickets</a><a href=\"/convocados\">Convocados</a><a href=\"/convocados/importar\">Importar padrón</a></nav></header>""",
)

MUNICIPIO_PREFIJO = {
    "Amealco de Bonfil": "AME", "Arroyo Seco": "ARS", "Cadereyta de Montes": "CAD",
    "Colón": "COL", "Corregidora": "COR", "El Marqués": "MAR",
    "Ezequiel Montes": "EZM", "Huimilpan": "HUI", "Jalpan de Serra": "JAL",
    "Landa de Matamoros": "LAM", "Pedro Escobedo": "PED", "Peñamiller": "PEN",
    "Pinal de Amoles": "PIN", "Querétaro": "QRO", "San Joaquín": "SAJ",
    "San Juan del Río": "SJR", "Tequisquiapan": "TEQ", "Tolimán": "TOL",
}

SCHEMA = """
CREATE TABLE IF NOT EXISTS convocados (
    id                   BIGSERIAL PRIMARY KEY,
    clave_externa        TEXT,
    codigo_verificacion  TEXT UNIQUE,
    nombre_completo      TEXT NOT NULL,
    municipio            TEXT NOT NULL,
    telefono             TEXT,
    cargo                TEXT,
    seccion              TEXT,
    correo               TEXT,
    estado_convocatoria  TEXT NOT NULL DEFAULT 'Convocado',
    observaciones        TEXT,
    telegram_id          BIGINT,
    telegram_username    TEXT,
    telegram_nombre      TEXT,
    telegram_verificado  BOOLEAN NOT NULL DEFAULT FALSE,
    fecha_verificacion   TIMESTAMPTZ,
    ultima_interaccion   TIMESTAMPTZ,
    estado_capacitacion  TEXT NOT NULL DEFAULT 'Pendiente',
    creado_en            TIMESTAMPTZ NOT NULL DEFAULT now(),
    actualizado_en       TIMESTAMPTZ NOT NULL DEFAULT now()
);
CREATE UNIQUE INDEX IF NOT EXISTS idx_convocados_telegram_unique
    ON convocados(telegram_id) WHERE telegram_id IS NOT NULL;
CREATE INDEX IF NOT EXISTS idx_convocados_municipio ON convocados(municipio);
CREATE INDEX IF NOT EXISTS idx_convocados_verificado ON convocados(telegram_verificado);
CREATE INDEX IF NOT EXISTS idx_convocados_codigo_upper ON convocados(upper(codigo_verificacion));
"""


def ensure_schema(conn=None):
    own = conn is None
    conn = conn or base.db()
    try:
        conn.execute(SCHEMA)
        if own:
            conn.commit()
    finally:
        if own:
            conn.close()


@app.on_event("startup")
def startup_convocados():
    ensure_schema()


def txt(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def norm_code(v: Any) -> str | None:
    s = txt(v)
    return "".join(s.upper().split()) if s else None


def normalizar_headers(headers):
    return [str(h).strip().upper() if h is not None else "" for h in headers]


def parse_xlsx(data: bytes) -> list[dict]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb["Convocados"] if "Convocados" in wb.sheetnames else wb.active
    header_row = None
    headers = None
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
        hs = normalizar_headers(row)
        if "NOMBRE_COMPLETO" in hs and "MUNICIPIO" in hs:
            header_row, headers = idx, hs
            break
    if not header_row:
        raise ValueError("No encontré los encabezados NOMBRE_COMPLETO y MUNICIPIO en las primeras 20 filas.")
    rows = []
    for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
        rec = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers)) if headers[i]}
        if not any(txt(v) for v in rec.values()):
            continue
        rows.append(rec)
    return rows


def parse_csv(data: bytes) -> list[dict]:
    text = data.decode("utf-8-sig")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    rows = []
    for row in reader:
        rec = {str(k).strip().upper(): v for k, v in row.items() if k}
        if any(txt(v) for v in rec.values()):
            rows.append(rec)
    return rows


def prefijo_municipio(municipio: str) -> str:
    return MUNICIPIO_PREFIJO.get(municipio, "QRO")


def buscar_existente(conn, rec: dict):
    codigo = norm_code(rec.get("CODIGO_VERIFICACION"))
    clave = txt(rec.get("CLAVE_EXTERNA"))
    nombre = txt(rec.get("NOMBRE_COMPLETO"))
    municipio = txt(rec.get("MUNICIPIO"))
    telefono = txt(rec.get("TELEFONO"))

    if codigo:
        row = conn.execute(
            "SELECT * FROM convocados WHERE upper(codigo_verificacion)=upper(%s)", (codigo,)
        ).fetchone()
        if row:
            return row
    if clave:
        row = conn.execute(
            "SELECT * FROM convocados WHERE lower(clave_externa)=lower(%s) ORDER BY id LIMIT 1", (clave,)
        ).fetchone()
        if row:
            return row
    if nombre and municipio:
        if telefono:
            row = conn.execute(
                """SELECT * FROM convocados
                     WHERE lower(nombre_completo)=lower(%s) AND municipio=%s AND COALESCE(telefono,'')=%s
                     ORDER BY id LIMIT 1""",
                (nombre, municipio, telefono),
            ).fetchone()
        else:
            row = conn.execute(
                """SELECT * FROM convocados
                     WHERE lower(nombre_completo)=lower(%s) AND municipio=%s
                       AND COALESCE(seccion,'')=COALESCE(%s,'') AND COALESCE(cargo,'')=COALESCE(%s,'')
                     ORDER BY id LIMIT 1""",
                (nombre, municipio, txt(rec.get("SECCION")), txt(rec.get("CARGO"))),
            ).fetchone()
        if row:
            return row
    return None


def guardar_registro(conn, rec: dict) -> str:
    nombre = txt(rec.get("NOMBRE_COMPLETO"))
    municipio = txt(rec.get("MUNICIPIO"))
    if not nombre or not municipio:
        raise ValueError("NOMBRE_COMPLETO y MUNICIPIO son obligatorios")

    datos = {
        "clave_externa": txt(rec.get("CLAVE_EXTERNA")),
        "nombre_completo": nombre,
        "municipio": municipio,
        "telefono": txt(rec.get("TELEFONO")),
        "cargo": txt(rec.get("CARGO")),
        "seccion": txt(rec.get("SECCION")),
        "correo": txt(rec.get("CORREO")),
        "estado_convocatoria": txt(rec.get("ESTADO_CONVOCATORIA")) or "Convocado",
        "observaciones": txt(rec.get("OBSERVACIONES")),
        "estado_capacitacion": txt(rec.get("ESTADO_CAPACITACION")) or "Pendiente",
    }
    existente = buscar_existente(conn, rec)
    if existente:
        conn.execute(
            """
            UPDATE convocados SET
                clave_externa=%(clave_externa)s,
                nombre_completo=%(nombre_completo)s,
                municipio=%(municipio)s,
                telefono=%(telefono)s,
                cargo=%(cargo)s,
                seccion=%(seccion)s,
                correo=%(correo)s,
                estado_convocatoria=%(estado_convocatoria)s,
                observaciones=%(observaciones)s,
                estado_capacitacion=%(estado_capacitacion)s,
                actualizado_en=now()
            WHERE id=%(id)s
            """,
            {**datos, "id": existente["id"]},
        )
        return "actualizado"

    codigo = norm_code(rec.get("CODIGO_VERIFICACION"))
    row = conn.execute(
        """
        INSERT INTO convocados
            (clave_externa,codigo_verificacion,nombre_completo,municipio,telefono,cargo,seccion,correo,
             estado_convocatoria,observaciones,estado_capacitacion)
        VALUES
            (%(clave_externa)s,%(codigo)s,%(nombre_completo)s,%(municipio)s,%(telefono)s,%(cargo)s,%(seccion)s,%(correo)s,
             %(estado_convocatoria)s,%(observaciones)s,%(estado_capacitacion)s)
        RETURNING id
        """,
        {**datos, "codigo": codigo},
    ).fetchone()
    if not codigo:
        codigo = f"{prefijo_municipio(municipio)}-{row['id']:05d}"
        conn.execute(
            "UPDATE convocados SET codigo_verificacion=%s WHERE id=%s",
            (codigo, row["id"]),
        )
    return "insertado"


def verif_badge(ok: bool) -> str:
    return '<span class="ok">✅ VERIFICADO</span>' if ok else '<span class="warn">⏳ PENDIENTE</span>'


@app.get("/convocados", response_class=HTMLResponse)
def convocados(municipio: str = "", verificacion: str = "", _: str = Depends(base.auth)):
    ensure_schema()
    where, params = [], []
    if municipio:
        where.append("municipio=%s"); params.append(municipio)
    if verificacion == "verificados":
        where.append("telegram_verificado=TRUE")
    elif verificacion == "pendientes":
        where.append("telegram_verificado=FALSE")
    sql = "SELECT * FROM convocados"
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY municipio,nombre_completo"

    with base.db() as conn:
        rows = conn.execute(sql, params).fetchall()
        general = conn.execute(
            """SELECT count(*) total,
                      count(*) FILTER (WHERE telegram_verificado) verificados,
                      count(*) FILTER (WHERE NOT telegram_verificado) pendientes
                 FROM convocados"""
        ).fetchone()
        municipios = conn.execute("SELECT DISTINCT municipio FROM convocados ORDER BY municipio").fetchall()

    total, ver, pen = general["total"], general["verificados"], general["pendientes"]
    pct = (ver / total * 100) if total else 0
    cards = (
        f'<div class="cards"><div class="card"><b>{total}</b>Convocados</div>'
        f'<div class="card"><b style="color:#16a34a">{ver}</b>Verificados</div>'
        f'<div class="card"><b style="color:#d97706">{pen}</b>Pendientes</div>'
        f'<div class="card"><b>{pct:.1f}%</b>Avance</div></div>'
    )
    opt_mun = '<option value="">Todos</option>' + "".join(
        f'<option value="{base.h(m["municipio"])}"{" selected" if m["municipio"]==municipio else ""}>{base.h(m["municipio"])}</option>'
        for m in municipios
    )
    opt_ver = "".join([
        f'<option value=""{" selected" if not verificacion else ""}>Todos</option>',
        f'<option value="verificados"{" selected" if verificacion=="verificados" else ""}>Verificados</option>',
        f'<option value="pendientes"{" selected" if verificacion=="pendientes" else ""}>Pendientes</option>',
    ])
    filtros = f"""<form class="filtros" method="get">
      <div><label>Municipio</label><select name="municipio">{opt_mun}</select></div>
      <div><label>Telegram</label><select name="verificacion">{opt_ver}</select></div>
      <button type="submit">Filtrar</button><a href="/convocados" style="align-self:center">Limpiar</a>
    </form>"""
    acciones = """<div class="actions">
      <a href="/convocados/importar">⬆️ Importar Excel/CSV</a>
      <a href="/convocados/export">⬇️ Exportar control</a>
    </div>"""

    body_rows = ""
    for r in rows:
        fecha = r["fecha_verificacion"].strftime("%d/%m/%Y %H:%M") if r.get("fecha_verificacion") else "—"
        usuario = f"@{base.h(r['telegram_username'])}" if r.get("telegram_username") else "—"
        body_rows += f"""<tr>
          <td><code>{base.h(r['codigo_verificacion'])}</code></td>
          <td>{base.h(r['nombre_completo'])}</td><td>{base.h(r['municipio'])}</td>
          <td>{base.h(r.get('cargo') or '—')}</td><td>{base.h(r.get('seccion') or '—')}</td>
          <td>{base.h(r.get('telefono') or '—')}</td><td>{verif_badge(r['telegram_verificado'])}</td>
          <td>{usuario}</td><td>{fecha}</td><td>{base.h(r.get('estado_capacitacion') or 'Pendiente')}</td>
        </tr>"""
    if not body_rows:
        body_rows = '<tr><td colspan="10" class="muted">No hay convocados con estos filtros.</td></tr>'
    tabla = f"""<table><thead><tr><th>Código</th><th>Nombre</th><th>Municipio</th><th>Cargo</th>
      <th>Sección</th><th>Teléfono</th><th>Telegram</th><th>Usuario</th><th>Fecha</th><th>Capacitación</th>
      </tr></thead><tbody>{body_rows}</tbody></table>"""
    return base.PAGE.format(body=f"<h2>Control de convocados</h2>{cards}{acciones}{filtros}{tabla}")


@app.get("/convocados/importar", response_class=HTMLResponse)
def importar_form(_: str = Depends(base.auth)):
    ensure_schema()
    body = """<p><a href="/convocados">← Volver a convocados</a></p>
    <div class="box"><h2>Importar padrón de convocados</h2>
    <p>Sube la plantilla <b>Excel (.xlsx)</b> o un <b>CSV</b>. Se requieren las columnas
    <code>NOMBRE_COMPLETO</code> y <code>MUNICIPIO</code>. Si el archivo no trae
    <code>CODIGO_VERIFICACION</code>, el sistema lo genera automáticamente.</p>
    <p class="muted">Las columnas de Telegram del Excel se ignoran por seguridad: solamente el bot puede marcar a una persona como verificada.</p>
    <form method="post" enctype="multipart/form-data">
      <div style="margin:18px 0"><input type="file" name="file" accept=".xlsx,.csv" required></div>
      <button type="submit">Importar padrón</button>
    </form></div>"""
    return base.PAGE.format(body=body)


@app.post("/convocados/importar", response_class=HTMLResponse)
async def importar_post(file: UploadFile = File(...), _: str = Depends(base.auth)):
    nombre = (file.filename or "").lower()
    data = await file.read()
    if len(data) > 10 * 1024 * 1024:
        raise HTTPException(413, "El archivo excede 10 MB")
    try:
        if nombre.endswith(".xlsx"):
            rows = parse_xlsx(data)
        elif nombre.endswith(".csv"):
            rows = parse_csv(data)
        else:
            raise ValueError("Formato no soportado. Usa .xlsx o .csv")
    except Exception as exc:
        return HTMLResponse(base.PAGE.format(body=f'<div class="box"><h2>Error de lectura</h2><p class="danger">{base.h(exc)}</p><p><a href="/convocados/importar">Volver</a></p></div>'), status_code=400)

    insertados = actualizados = omitidos = 0
    errores = []
    ensure_schema()
    with base.db() as conn:
        for n, rec in enumerate(rows, start=1):
            try:
                resultado = guardar_registro(conn, rec)
                if resultado == "insertado": insertados += 1
                else: actualizados += 1
            except Exception as exc:
                omitidos += 1
                errores.append(f"Fila {n}: {exc}")
        conn.commit()

    err_html = ""
    if errores:
        err_html = "<h3>Observaciones</h3><ul>" + "".join(f"<li>{base.h(e)}</li>" for e in errores[:30]) + "</ul>"
    body = f"""<div class="box"><h2>Importación terminada</h2>
      <p class="ok">Insertados: <b>{insertados}</b></p>
      <p>Actualizados: <b>{actualizados}</b></p>
      <p>Omitidos con error: <b>{omitidos}</b></p>{err_html}
      <p><a href="/convocados">Ver control de convocados →</a></p></div>"""
    return base.PAGE.format(body=body)


@app.get("/convocados/export")
def export_convocados(_: str = Depends(base.auth)):
    ensure_schema()
    with base.db() as conn:
        rows = conn.execute("SELECT * FROM convocados ORDER BY municipio,nombre_completo").fetchall()
    buf = io.StringIO()
    if rows:
        campos = [
            "clave_externa", "codigo_verificacion", "nombre_completo", "municipio", "telefono",
            "cargo", "seccion", "correo", "estado_convocatoria", "observaciones",
            "telegram_verificado", "fecha_verificacion", "telegram_id", "telegram_username",
            "estado_capacitacion"
        ]
        w = csv.DictWriter(buf, fieldnames=campos)
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k) for k in campos})
    data = buf.getvalue().encode("utf-8-sig")
    return StreamingResponse(
        io.BytesIO(data), media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=control_convocados_telegram.csv"},
    )
